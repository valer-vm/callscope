"""«Звонки»: список с фильтрами, страница звонка, правки, смена сторон, удаление, Excel."""

from __future__ import annotations

import io
import re
from datetime import date, timedelta
from urllib.parse import urlencode

import pytest
from openpyxl import load_workbook

from callscope import pipeline, stats
from callscope.checklist import load_checklist
from callscope.config import load_config
from callscope.storage import Storage


def _stamp(days_ago: int) -> str:
    return f"{date.today() - timedelta(days=days_ago):%Y-%m-%d}_1030"


def _add(storage: Storage, raw_from_fixture, normalizer, fixture: str, name: str) -> int:
    record, _ = storage.add_upload(name, (fixture + name).encode())
    storage.next_queued()
    config = load_config(storage.config_path)
    checklist = load_checklist(storage.checklist_path)
    raw = raw_from_fixture(fixture)
    result = pipeline.analyze_raw(name, raw, config, checklist, normalizer)
    storage.save_result(record.id, raw, result)
    return record.id


@pytest.fixture
def storage(web_settings) -> Storage:
    return Storage(web_settings.data_dir)


@pytest.fixture
def seeded(storage, raw_from_fixture, normalizer) -> dict[str, int]:
    ids = {
        "good": _add(storage, raw_from_fixture, normalizer, "good.json",
                     f"ivanov_{_stamp(3)}.m4a"),
        "bad": _add(storage, raw_from_fixture, normalizer, "no_need.json",
                    f"petrova_{_stamp(10)}.m4a"),
        "old": _add(storage, raw_from_fixture, normalizer, "no_need.json",
                    f"petrova_{_stamp(200)}.m4a"),
    }
    ids["queued"] = storage.add_upload("queued.mp3", b"q")[0].id
    return ids


def _row_ids(html: str) -> list[int]:
    return [int(m) for m in re.findall(r'<a href="/calls/(\d+)"', html)]


# --- список ------------------------------------------------------------------------------


def test_pages_require_login(client, seeded):
    assert client.get("/calls").status_code == 303
    assert client.get(f"/calls/{seeded['good']}").status_code == 303
    assert client.get("/calls/export.xlsx").status_code == 401
    for suffix in ("correction", "swap", "delete"):
        assert client.post(f"/calls/{seeded['good']}/{suffix}").status_code == 303


def test_empty_list_shows_hint(authed):
    html = authed.get("/calls").text
    assert "Звонков пока нет" in html or "Звонков не найдено" in html
    assert 'href="/upload"' in html


def test_list_shows_done_calls_in_default_period(authed, seeded):
    html = authed.get("/calls").text
    assert set(_row_ids(html)) == {seeded["good"], seeded["bad"]}  # old вне 4 недель
    assert "petrova_" in html and "ivanov_" in html
    assert "Ещё не обработаны" in html  # звонок в очереди не мешает, но упомянут


def test_filters_live_in_the_url(authed, seeded):
    # Фильтры независимы от сессии: период в адресе.
    url = "/calls?period=all"
    assert set(_row_ids(authed.get(url).text)) == {seeded[k] for k in ("good", "bad", "old")}
    manager = Storage(authed.app.state.settings.data_dir).get(seeded["bad"]).manager
    by_manager = authed.get(url + "&" + urlencode({"manager": manager})).text
    assert set(_row_ids(by_manager)) == {seeded["bad"], seeded["old"]}
    only_good = authed.get(url + "&score_min=90").text
    assert _row_ids(only_good) == [seeded["good"]]
    only_bad = authed.get(url + "&score_max=50").text
    assert set(_row_ids(only_bad)) == {seeded["bad"], seeded["old"]}


def test_filter_by_failed_item(authed, seeded, storage):
    call = storage.get(seeded["bad"])
    failed = next(i.item_id for i in call.score.items if i.status == "failed")
    html = authed.get(f"/calls?period=all&failed_item={failed}").text
    assert seeded["good"] not in _row_ids(html)
    assert seeded["bad"] in _row_ids(html)


def test_custom_period_dates(authed, seeded):
    start = date.today() - timedelta(days=5)
    html = authed.get(f"/calls?period=custom&start={start}&end={date.today()}").text
    assert _row_ids(html) == [seeded["good"]]


def test_pagination_fifty_per_page(authed, storage, raw_from_fixture, normalizer):
    config = load_config(storage.config_path)
    checklist = load_checklist(storage.checklist_path)
    raw = raw_from_fixture("good.json")
    result = pipeline.analyze_raw("x.m4a", raw, config, checklist, normalizer)
    for n in range(55):
        record, _ = storage.add_upload(f"ivanov_{_stamp(1)}_{n}.m4a", f"c{n}".encode())
        storage.next_queued()
        storage.save_result(record.id, raw, result)
    first = _row_ids(authed.get("/calls?period=all").text)
    second = _row_ids(authed.get("/calls?period=all&page=2").text)
    assert len(first) == 50 and len(second) == 5
    assert not set(first) & set(second)
    assert authed.get("/calls?period=all&page=99").status_code == 200


def test_garbage_filters_do_not_break_page(authed, seeded):
    response = authed.get("/calls?score_min=abc&start=zzz&period=nope&page=-3")
    assert response.status_code == 200


def test_list_marks_manual_corrections(authed, seeded, storage):
    call = storage.get(seeded["bad"])
    failed = next(i.item_id for i in call.score.items if i.status == "failed")
    assert "исправлено вручную" not in authed.get("/calls").text
    storage.set_correction(seeded["bad"], failed, "passed", "")
    assert "исправлено вручную" in authed.get("/calls").text


# --- страница звонка -----------------------------------------------------------------------


def test_call_page_has_player_and_timestamped_elements(authed, seeded, storage):
    call = storage.get(seeded["good"])
    html = authed.get(f"/calls/{seeded['good']}").text
    assert f'src="/audio/{seeded["good"]}"' in html and "<audio" in html
    assert "Менеджер" in html and "Клиент" in html
    for utterance in call.dialogue.utterances:
        assert f'data-start="{utterance.start:.2f}"' in html
    # у цитат тоже есть data-start с временем реплики-доказательства
    quotes = [e for item in call.score.items for e in item.evidence]
    assert quotes
    for evidence in quotes:
        assert f'class="quote" data-start="{evidence.start:.2f}"' in html
    for item in call.score.items:
        assert item.reason in html


def test_call_page_mono_marks_approximate(authed, storage, raw_from_fixture, normalizer):
    from dataclasses import replace

    raw = replace(raw_from_fixture("good.json"), mode="mono")
    config = load_config(storage.config_path)
    checklist = load_checklist(storage.checklist_path)
    result = pipeline.analyze_raw("ivanov_2026-10-01_1000.m4a", raw, config, checklist, normalizer)
    record, _ = storage.add_upload("ivanov_2026-10-01_1000.m4a", b"mono")
    storage.next_queued()
    storage.save_result(record.id, raw, result)
    assert "стороны определены приблизительно" in authed.get(f"/calls/{record.id}").text


def _wav(seconds: float) -> bytes:
    import wave

    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(1)
        handle.setframerate(8000)
        handle.writeframes(b"\x80" * int(8000 * seconds))
    return buffer.getvalue()


def _add_wav(storage, raw_from_fixture, normalizer, seconds: float) -> int:
    raw = raw_from_fixture("good.json")
    config = load_config(storage.config_path)
    checklist = load_checklist(storage.checklist_path)
    name = "ivanov_2026-10-01_1000.wav"
    result = pipeline.analyze_raw(name, raw, config, checklist, normalizer)
    record, _ = storage.add_upload(name, _wav(seconds))
    storage.next_queued()
    storage.save_result(record.id, raw, result)
    return record.id


def test_call_without_real_audio_shows_note_instead_of_player(
    authed, storage, raw_from_fixture, normalizer
):
    # Демо-заглушка: секунда тишины при расшифровке на десятки секунд.
    call_id = _add_wav(storage, raw_from_fixture, normalizer, seconds=1.0)
    html = authed.get(f"/calls/{call_id}").text
    assert "<audio" not in html
    assert "Записи разговора нет" in html


def test_call_with_full_wav_keeps_player(authed, storage, raw_from_fixture, normalizer):
    duration = raw_from_fixture("good.json").duration
    call_id = _add_wav(storage, raw_from_fixture, normalizer, seconds=duration)
    html = authed.get(f"/calls/{call_id}").text
    assert f'src="/audio/{call_id}"' in html
    assert "Записи разговора нет" not in html


def test_call_page_unknown_is_404(authed):
    assert authed.get("/calls/12345").status_code == 404


def test_call_page_not_ready_states_status(authed, seeded):
    html = authed.get(f"/calls/{seeded['queued']}").text
    assert "в очереди" in html
    assert "<audio" not in html


def test_call_page_error_status_shows_reason(authed, storage):
    record, _ = storage.add_upload("broken.mp3", b"e")
    storage.next_queued()
    storage.save_error(record.id, "Файл повреждён, прочитать нечего.")
    html = authed.get(f"/calls/{record.id}").text
    assert "Файл повреждён" in html


# --- правки --------------------------------------------------------------------------------


def _failed_item(storage, call_id) -> str:
    return next(i.item_id for i in storage.get(call_id).score.items if i.status == "failed")


def test_correction_changes_score_and_overview(authed, seeded, storage, web_settings):
    call_id = seeded["bad"]
    before = storage.get(call_id).score.score
    item = _failed_item(storage, call_id)
    response = authed.post(
        f"/calls/{call_id}/correction",
        data={"item_id": item, "status": "passed", "comment": "Сказал в конце разговора"},
    )
    assert response.status_code == 303
    after = storage.get(call_id)
    assert after.score.score > before
    assert after.corrections[item].comment == "Сказал в конце разговора"
    html = authed.get(f"/calls/{call_id}").text
    assert "исправлено вручную" in html
    assert "Сказал в конце разговора" in html
    assert f">{after.score.score:.0f}<" in html  # новый балл в шапке страницы
    # обзор считает по правке: провалов пункта стало на один меньше
    cfg = load_config(storage.config_path)
    chk = load_checklist(storage.checklist_path)
    period = stats.period_preset("all", date.today())

    def failed_count() -> int:
        data = stats.overview(storage.list_calls(stats.CallFilters(period)), chk,
                              cfg.thresholds, period)
        return next(s for s in data.funnel if s.item_id == item).failed

    assert failed_count() == 1  # остался только звонок «old»
    storage.clear_correction(call_id, item)
    assert failed_count() == 2


def test_comment_only_keeps_auto_status(authed, seeded, storage):
    call_id = seeded["bad"]
    item = _failed_item(storage, call_id)
    before = storage.get(call_id).score.score
    authed.post(f"/calls/{call_id}/correction",
                data={"item_id": item, "status": "auto", "comment": "Проверить"})
    record = storage.get(call_id)
    assert record.score.score == before
    assert record.corrections[item].status is None
    assert "Проверить" in authed.get(f"/calls/{call_id}").text


def test_clear_correction_restores_auto(authed, seeded, storage):
    call_id = seeded["bad"]
    item = _failed_item(storage, call_id)
    before = storage.get(call_id).score.score
    authed.post(f"/calls/{call_id}/correction",
                data={"item_id": item, "status": "passed", "comment": "x"})
    response = authed.post(f"/calls/{call_id}/correction",
                           data={"item_id": item, "action": "clear"})
    assert response.status_code == 303
    record = storage.get(call_id)
    assert record.score.score == before and not record.corrections


def test_correction_rejects_bad_input(authed, seeded):
    call_id = seeded["bad"]
    url = f"/calls/{call_id}/correction"
    assert authed.post(url, data={"item_id": "no_such", "status": "passed"}).status_code == 400
    assert authed.post(url, data={"item_id": "greeting", "status": "weird"}).status_code == 400
    assert authed.post("/calls/999/correction",
                       data={"item_id": "greeting", "status": "passed"}).status_code == 404


# --- стороны и удаление ----------------------------------------------------------------------


def test_swap_sides_flips_roles_and_keeps_corrections(authed, seeded, storage):
    call_id = seeded["bad"]
    item = _failed_item(storage, call_id)
    authed.post(f"/calls/{call_id}/correction",
                data={"item_id": item, "status": "na", "comment": "ручная"})
    before = storage.get(call_id).dialogue
    response = authed.post(f"/calls/{call_id}/swap")
    assert response.status_code == 303
    after = storage.get(call_id)
    assert after.dialogue.manager_channel != before.manager_channel
    assert after.corrections[item].comment == "ручная"


def test_delete_removes_call_and_audio(authed, seeded, storage):
    call_id = seeded["good"]
    audio_path = storage.get(call_id).audio_path
    assert audio_path.exists()
    response = authed.post(f"/calls/{call_id}/delete")
    assert response.status_code == 303 and response.headers["location"] == "/calls"
    assert call_id not in _row_ids(authed.get("/calls?period=all").text)
    assert authed.get(f"/calls/{call_id}").status_code == 404
    assert authed.get(f"/audio/{call_id}").status_code == 404
    assert not audio_path.exists()


# --- Excel ----------------------------------------------------------------------------------


def test_export_has_filtered_calls_with_corrections(authed, seeded, storage):
    item = _failed_item(storage, seeded["bad"])
    authed.post(f"/calls/{seeded['bad']}/correction",
                data={"item_id": item, "status": "passed", "comment": ""})
    manager = storage.get(seeded["bad"]).manager
    response = authed.get("/calls/export.xlsx?" + urlencode({"period": "all", "manager": manager}))
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert f"callscope_{date.today():%Y-%m-%d}.xlsx" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert len(book.sheetnames) == 5
    rows = list(book["Звонки"].iter_rows(min_row=2, values_only=True))
    names = {row[0] for row in rows}
    assert names <= {storage.get(seeded[k]).file_name for k in ("bad", "old")}
    assert names  # фильтр по менеджеру не пустой
    assert storage.get(seeded["good"]).file_name not in names
    corrected = next(r for r in rows if r[0] == storage.get(seeded["bad"]).file_name)
    assert corrected[3] == storage.get(seeded["bad"]).score.score
    assert book["Расшифровки"].max_row > 1  # реплики в файле есть


def test_export_follows_same_filters_as_list(authed, seeded):
    response = authed.get("/calls/export.xlsx?period=all&score_min=90")
    book = load_workbook(io.BytesIO(response.content))
    assert len(list(book["Звонки"].iter_rows(min_row=2, values_only=True))) == 1


def test_export_empty_selection_still_valid_file(authed):
    response = authed.get("/calls/export.xlsx?period=all")
    assert response.status_code == 200
    assert len(load_workbook(io.BytesIO(response.content)).sheetnames) == 5
