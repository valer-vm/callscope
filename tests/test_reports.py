from datetime import datetime

from openpyxl import load_workbook

from callscope.aggregate import build_summary
from callscope.checklist import Checklist, ChecklistItem
from callscope.config import Thresholds
from callscope.model import (
    CallMeta,
    CallResult,
    CallScore,
    Dialogue,
    Evidence,
    ItemResult,
    SkippedFile,
    Utterance,
)
from callscope.report_html import write_html
from callscope.report_xlsx import write_xlsx

CHECKLIST = Checklist(
    items=(
        ChecklistItem(id="greeting", title="Приветствие и представление", weight=5),
        ChecklistItem(id="need", title="Выявление потребности", weight=20),
    )
)
EVIL = "<script>alert(1)</script>"


def make_call(name: str, manager: str, score: float, grade: str, approximate: bool = False):
    quote = Evidence(0, 65.0, "manager", "Здравствуйте, меня зовут Анна")
    items = (
        ItemResult("greeting", "passed", "Приветствие найдено", (quote,)),
        ItemResult(
            "need", "failed", "Менеджер не задал открытых вопросов",
            (Evidence(1, 12.0, "client", EVIL),),
        ),
    )
    dialogue = Dialogue(
        utterances=(
            Utterance(0, 65.0, 70.0, "Здравствуйте, меня зовут Анна", "manager"),
            Utterance(1, 12.0, 15.0, EVIL, "client"),
        ),
        approximate=approximate,
        manager_channel="left",
        role_source="intro",
    )
    return CallResult(
        CallMeta(name, manager, datetime(2026, 10, 12, 15, 30), True),
        dialogue,
        CallScore(items, score, "need", grade),
        ("Тестовое предупреждение",),
    )


def summary():
    calls = [
        make_call("anna_1.wav", "anna", 90, "good"),
        make_call("boris_1.wav", "boris", 30, "bad", approximate=True),
    ]
    return build_summary(
        calls, [SkippedFile("broken.mp3", "файл не читается")], CHECKLIST, Thresholds(), "calls",
        generated_at=datetime(2026, 10, 16, 12, 0),
    )


def test_html_contains_key_content(tmp_path):
    path = tmp_path / "out" / "report.html"
    write_html(summary(), path)
    html = path.read_text(encoding="utf-8")
    assert summary().conclusion in html
    expected = (
        "anna_1.wav", "boris", "Выявление потребности",
        "Менеджер не задал открытых вопросов", "Здравствуйте, меня зовут Анна",
        "Стороны определены приблизительно", "broken.mp3", "файл не читается",
        "Тестовое предупреждение", "01:05",
    )
    for text in expected:
        assert text in html
    assert "score good" in html and "score bad" in html
    assert "@media print" in html


def test_html_is_self_contained(tmp_path):
    path = tmp_path / "report.html"
    write_html(summary(), path)
    html = path.read_text(encoding="utf-8")
    assert "http://" not in html and "https://" not in html
    assert "<link" not in html and "src=" not in html


def test_html_escapes_transcript_text(tmp_path):
    path = tmp_path / "report.html"
    write_html(summary(), path)
    html = path.read_text(encoding="utf-8")
    assert EVIL not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_xlsx_sheets_and_content(tmp_path):
    path = tmp_path / "report.xlsx"
    write_xlsx(summary(), path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Сводка", "Менеджеры", "Звонки", "Пункты", "Расшифровки"]

    assert wb["Сводка"]["B1"].value == summary().conclusion
    managers = [[c.value for c in row] for row in wb["Менеджеры"].iter_rows(min_row=2)]
    assert managers[0][:3] == ["anna", 1, 90.0]
    assert wb["Менеджеры"]["A1"].font.bold

    calls = [[c.value for c in row] for row in wb["Звонки"].iter_rows(min_row=2)]
    assert calls[1][0] == "boris_1.wav" and calls[1][3] == 30.0
    assert calls[1][5] == "Выявление потребности" and calls[1][6] == "приблизительно"

    items = [[c.value for c in row] for row in wb["Пункты"].iter_rows(min_row=2)]
    assert len(items) == 4
    failed = next(r for r in items if r[3] == "Провален")
    assert failed[2] == "Выявление потребности"
    assert "Менеджер не задал открытых вопросов" in failed[4]
    assert "00:12 Клиент" in failed[5]

    transcripts = [[c.value for c in row] for row in wb["Расшифровки"].iter_rows(min_row=2)]
    first_line = ["anna_1.wav", "anna", "01:05", "Менеджер", "Здравствуйте, меня зовут Анна"]
    assert first_line in transcripts


def test_xlsx_colors_score_and_widths(tmp_path):
    path = tmp_path / "report.xlsx"
    write_xlsx(summary(), path)
    ws = load_workbook(path)["Звонки"]
    assert ws["D2"].fill.fgColor.rgb.endswith("DCFCE7")
    assert ws["D3"].fill.fgColor.rgb.endswith("FEE2E2")
    assert ws.column_dimensions["A"].width >= 20


def test_reports_for_summary_without_calls(tmp_path):
    s = build_summary([], [SkippedFile("x.mp3", "не читается")], CHECKLIST, Thresholds(), "calls")
    write_html(s, tmp_path / "e.html")
    write_xlsx(s, tmp_path / "e.xlsx")
    assert "x.mp3" in (tmp_path / "e.html").read_text(encoding="utf-8")
    assert load_workbook(tmp_path / "e.xlsx").sheetnames[0] == "Сводка"


def test_xlsx_writes_equals_sign_text_as_text_not_formula(tmp_path):
    formula_like = '=HYPERLINK("http://evil","клик")'
    call = make_call("anna_1.wav", "=anna", 90, "good")
    dialogue = Dialogue(
        utterances=(Utterance(0, 0.0, 2.0, formula_like, "manager"),),
        approximate=False,
        manager_channel="left",
        role_source="intro",
    )
    call = CallResult(call.meta, dialogue, call.score, ("=1+1",))
    data = build_summary([call], [], CHECKLIST, Thresholds(), "calls")
    path = tmp_path / "report.xlsx"
    write_xlsx(data, path)
    wb = load_workbook(path)
    cells = [cell for ws in wb for row in ws.iter_rows() for cell in row]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    texts = {c.value for c in cells}
    assert {formula_like, "=anna", "=1+1"} <= texts
