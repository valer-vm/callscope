"""Excel-отчёт из пяти листов."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from callscope.aggregate import Summary
from callscope.report_common import (
    GRADE_LABEL,
    ROLE_LABEL,
    STATUS_LABEL,
    fmt_datetime,
    fmt_time,
    titles,
)
from callscope.scoring import grade_for

_HEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
_GRADE_FILL = {
    "good": PatternFill("solid", fgColor="DCFCE7"),
    "medium": PatternFill("solid", fgColor="FEF3C7"),
    "bad": PatternFill("solid", fgColor="FEE2E2"),
}
_STATUS_FILL = {
    "passed": _GRADE_FILL["good"],
    "failed": _GRADE_FILL["bad"],
    "na": PatternFill("solid", fgColor="F3F4F6"),
}
_WRAP = Alignment(wrap_text=True, vertical="top")


def _table(ws: Worksheet, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
    for row in rows:
        ws.append(row)
    for letter_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=letter_idx).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP
    ws.freeze_panes = "A2"


def write_xlsx(summary: Summary, path: Path) -> None:
    names = titles(summary)
    thr = summary.thresholds
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 80
    ws.append(["Вывод", summary.conclusion])
    ws.append(["Папка", summary.source_dir])
    ws.append(["Сформирован", summary.generated_at.strftime("%d.%m.%Y %H:%M")])
    ws.append(["Звонков оценено", len(summary.calls)])
    ws.append(["Средний балл", round(summary.avg_score, 1)])
    ws.append(["Пороги", f"хороший ≥ {thr.good:g}, средний ≥ {thr.medium:g}, иначе плохой"])
    ws.append(["Пропущено файлов", len(summary.skipped)])
    ws.append([])
    ws.append(["Этап скрипта", "Потеряно звонков"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for item_id, n in summary.loss_distribution:
        ws.append([names.get(item_id, item_id), n])
    ws.append([])
    ws.append(["Рейтинг провалов: пункт", "Раз провален"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for item_id, n in summary.failure_ranking:
        ws.append([names.get(item_id, item_id), n])
    if summary.skipped:
        ws.append([])
        ws.append(["Пропущенный файл", "Причина"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for f in summary.skipped:
            ws.append([f.file_name, f.reason])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = _WRAP
    for r in range(1, 8):
        ws.cell(row=r, column=1).font = Font(bold=True)

    ws = wb.create_sheet("Менеджеры")
    rows = []
    for m in summary.managers:
        fails = "; ".join(f"{names.get(i, i)} ({n})" for i, n in m.top_failures) or "провалов нет"
        rows.append([m.manager, m.calls, round(m.avg_score, 1), fails])
    header = ["Менеджер", "Звонков", "Средний балл", "Чаще всего проваливает"]
    _table(ws, header, rows, [24, 10, 14, 70])
    for row_idx, m in enumerate(summary.managers, start=2):
        ws.cell(row=row_idx, column=3).fill = _GRADE_FILL[grade_for(m.avg_score, thr)]

    ws = wb.create_sheet("Звонки")
    rows = [
        [
            c.meta.file_name,
            c.meta.manager,
            fmt_datetime(c.meta.call_datetime),
            round(c.score.score, 1),
            GRADE_LABEL[c.score.grade],
            names.get(c.score.loss_item_id, c.score.loss_item_id)
            if c.score.loss_item_id
            else "нет",
            "приблизительно" if c.dialogue.approximate else "точно",
            "; ".join(c.warnings),
        ]
        for c in summary.calls
    ]
    header = [
        "Файл", "Менеджер", "Дата и время", "Балл", "Оценка", "Этап потери", "Стороны",
        "Предупреждения",
    ]
    _table(ws, header, rows, [30, 20, 18, 8, 12, 36, 16, 50])
    for row_idx, c in enumerate(summary.calls, start=2):
        ws.cell(row=row_idx, column=4).fill = _GRADE_FILL[c.score.grade]
        ws.cell(row=row_idx, column=5).fill = _GRADE_FILL[c.score.grade]

    ws = wb.create_sheet("Пункты")
    rows, statuses = [], []
    for c in summary.calls:
        for r in c.score.items:
            quote = "\n".join(
                f"{fmt_time(e.start)} {ROLE_LABEL[e.role]}: «{e.text}»" for e in r.evidence
            )
            reason = r.reason + (" (подтверждено моделью)" if r.by_model else "")
            title = names.get(r.item_id, r.item_id)
            rows.append(
                [c.meta.file_name, c.meta.manager, title, STATUS_LABEL[r.status], reason, quote]
            )
            statuses.append(r.status)
    header = ["Файл", "Менеджер", "Пункт", "Статус", "Причина", "Цитата"]
    _table(ws, header, rows, [30, 20, 36, 14, 60, 70])
    for row_idx, status in enumerate(statuses, start=2):
        ws.cell(row=row_idx, column=4).fill = _STATUS_FILL[status]

    ws = wb.create_sheet("Расшифровки")
    rows = [
        [c.meta.file_name, c.meta.manager, fmt_time(u.start), ROLE_LABEL[u.role], u.text]
        for c in summary.calls
        for u in c.dialogue.utterances
    ]
    _table(ws, ["Файл", "Менеджер", "Время", "Говорит", "Реплика"], rows, [30, 20, 9, 12, 100])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
